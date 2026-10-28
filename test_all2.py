# -*- coding: utf-8 -*-
import os, django, glob
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")
django.setup()
ROOT = r"C:\Users\acer\Desktop\test_chain\russia"
class C:
    def __init__(self, v): self.value = v
class GridSheet:
    def __init__(self, g): self._g=g; self.max_row=len(g); self.max_column=max((len(r) for r in g),default=0)
    def cell(self, row, column):
        r,c=row-1,column-1
        if 0<=r<len(self._g) and 0<=c<len(self._g[r]): return C(self._g[r][c])
        return C(None)
def read_grid(path):
    try:
        import xlrd; wb=xlrd.open_workbook(path); sh=wb.sheet_by_index(0)
        return [[sh.cell_value(r,c) for c in range(sh.ncols)] for r in range(sh.nrows)]
    except Exception:
        import openpyxl; wb=openpyxl.load_workbook(path,data_only=True,read_only=True); ws=wb.active
        return [list(r) for r in ws.iter_rows(values_only=True)]
from sales.pharmacy_all_configs import get_config, parse_chain
def nz(v):
    try: return float(str(v).replace(",",".")) != 0
    except: return bool(str(v).strip())
for path in sorted(glob.glob(os.path.join(ROOT,"*.xls*"))):
    fn=os.path.basename(path)
    if "(test)" in fn or fn.endswith("test.xls"): continue
    chain=fn.rsplit(".",1)[0]
    cfg,typ=get_config(chain)
    if cfg is None and typ not in ("avromed","kg","matrix","c1"):
        print(f"{chain:40} -> YOK"); continue
    try:
        sh=GridSheet(read_grid(path))
        rows=parse_chain(sh,chain,"",sh.max_row,sh.max_column,"Russia")
        total=len(rows) if rows else 0
        nzrows=[r for r in (rows or []) if nz(r.get("COUNT"))]
        s=""
        if nzrows:
            r0=nzrows[0]
            s=f" | or:{str(r0.get('PRODUCT',''))[:15]!r} ap:{str(r0.get('PHARMACY',''))[:13]!r} cnt:{r0.get('COUNT')} amt:{r0.get('AMOUNT')}"
        print(f"{chain:40} -> {typ or '?':7} top:{total} dolu:{len(nzrows)}{s}")
    except Exception as e:
        print(f"{chain:40} -> {typ or '?':7} HATA: {type(e).__name__} {str(e)[:35]}")