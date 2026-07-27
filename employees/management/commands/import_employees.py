"""
Management command: import employees from Olga's access-list spreadsheet.

Usage:
    python manage.py import_employees data/List_AccsessesHHHH.xlsx
"""

from datetime import date, datetime
from pathlib import Path
from typing import Any

from django.core.management.base import BaseCommand, CommandError

from employees.models import Employee

COL_NAME = 1
COL_UNIT = 2
COL_TITLE = 3
COL_EMAIL = 4
COL_ACTIVATION = 5
COL_REGION = 15
FIRST_DATA_ROW = 4


class Command(BaseCommand):
    """Import employees from the access-list .xlsx file."""

    help = "Import employees from Olga's access-list spreadsheet."

    def add_arguments(self, parser: Any) -> None:
        """Register the file-path argument."""
        parser.add_argument("path", type=str, help="Path to the .xlsx file")

    def handle(self, *args: Any, **options: Any) -> None:
        """Read the spreadsheet and upsert employee rows."""
        import openpyxl

        path = Path(options["path"])
        if not path.exists():
            raise CommandError(f"File not found: {path}")

        workbook = openpyxl.load_workbook(path, data_only=True)
        sheet = workbook.active

        created = 0
        updated = 0

        for row in range(FIRST_DATA_ROW, sheet.max_row + 1):
            name = self._clean(sheet.cell(row=row, column=COL_NAME).value)
            if not name:
                continue

            unit = self._clean(sheet.cell(row=row, column=COL_UNIT).value)
            title = self._clean(sheet.cell(row=row, column=COL_TITLE).value)
            email = self._clean(sheet.cell(row=row, column=COL_EMAIL).value)
            region = self._clean(sheet.cell(row=row, column=COL_REGION).value)
            activation = self._parse_date(
                sheet.cell(row=row, column=COL_ACTIVATION).value
            )

            defaults = {
                "full_name": name,
                "unit": unit,
                "title": title,
                "region": region,
                "activation_date": activation,
                "is_active": True,
            }

            if email:
                obj, was_created = Employee.objects.update_or_create(
                    email=email, defaults=defaults
                )
            else:
                obj, was_created = Employee.objects.update_or_create(
                    full_name=name, email="", defaults=defaults
                )

            if was_created:
                created += 1
            else:
                updated += 1

        self.stdout.write(
            self.style.SUCCESS(f"Done. Created: {created}, Updated: {updated}")
        )

    @staticmethod
    def _clean(value: Any) -> str:
        """Return a trimmed string, or empty string for None."""
        if value is None:
            return ""
        return str(value).strip()

    @staticmethod
    def _parse_date(value: Any) -> date | None:
        """Convert a spreadsheet cell into a date, or None if not parseable."""
        if value is None:
            return None
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%m/%d/%Y"):
            try:
                return datetime.strptime(str(value).strip(), fmt).date()
            except ValueError:
                continue
        return None